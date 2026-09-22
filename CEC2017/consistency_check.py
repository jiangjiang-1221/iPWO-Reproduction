# -*- coding: utf-8 -*-
"""
CEC2017 Dim10/30/50 reproduction data vs original folder consistency check
=================================================
For each reproduction folder CEC2017Dim{N}_TIFF_MATLAB/data/ and its original CEC2017Dim{N}/, compare item by item:
  (1) copy completeness: every *_Best_Values.xlsx / *_Avg_Convergence.xlsx exists and algorithm columns match
  (2) cell-level equality: except the SFOA column (col3) row 1 of Avg_Convergence, all other cells must equal the original file value by value
  (3) SFOA correction correctness: changed cell = row 2 value of the same-table SFOA column in the original (equivalent to curve(1)=curve(2))
  (4) SFOA no longer has a leading 0
  (5) nfe (x-axis end): Avg_Convergence rows - 1 matches the dimension expectation (Dim10=3335, Dim30=10001, Dim50=16668)
  (6) boxplot value consistency: reproduction data Best_Values and original folder Best_Values are equal cell by cell (boxplot is drawn directly from it)
Output PASS/FAIL summary; any FAIL lists details.
"""
import openpyxl, os, glob, numpy as np

ROOT = r"D:\File\GPTTest\CEC2017_Experiment\iPWO_Reproducible\CEC2017"
DIMS = [10, 30, 50]
EXPECT_NFE = {10: 3335, 30: 10001, 50: 16668}   # Avg_Convergence rows (excl header)

def load_matrix(fp):
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    header = list(rows[0])
    data = np.array(rows[1:], dtype=float)
    return header, data

def check_dim(dim):
    src = os.path.join(ROOT, f"CEC2017Dim{dim}")
    rep = os.path.join(ROOT, f"CEC2017Dim{dim}_TIFF_MATLAB", "data")
    problems = []
    funcs = [f for f in list(range(1,31)) if f != 2]   # CEC2017 has no F2

    for fid in funcs:
        bv = f"CEC2017_F{fid}_Dim{dim}_Best_Values.xlsx"
        av = f"CEC2017_F{fid}_Dim{dim}_Avg_Convergence.xlsx"
        src_bv, src_av = os.path.join(src, bv), os.path.join(src, av)
        rep_bv, rep_av = os.path.join(rep, bv), os.path.join(rep, av)

        # (1) copy completeness
        for p in [rep_bv, rep_av]:
            if not os.path.exists(p):
                problems.append(f"F{fid}: missing {os.path.basename(p)}")
                continue

        # Best_Values: equal to original value by value
        if os.path.exists(src_bv) and os.path.exists(rep_bv):
            h1, d1 = load_matrix(src_bv)
            h2, d2 = load_matrix(rep_bv)
            if list(h1) != list(h2):
                problems.append(f"F{fid} Best_Values header mismatch: {h1} vs {h2}")
            if d1.shape != d2.shape:
                problems.append(f"F{fid} Best_Values shape mismatch: {d1.shape} vs {d2.shape}")
            else:
                if not np.allclose(d1, d2, rtol=0, atol=1e-12, equal_nan=True):
                    diff = np.abs(d1 - d2)
                    problems.append(f"F{fid} Best_Values has {int((diff>1e-12).sum())} cells differing from original (max|Δ|={diff.max():.3e})")

        # Avg_Convergence: all equal to original except SFOA (col3) row 1
        if os.path.exists(src_av) and os.path.exists(rep_av):
            h1, d1 = load_matrix(src_av)   # original
            h2, d2 = load_matrix(rep_av)   # reproduction
            if list(h1) != list(h2):
                problems.append(f"F{fid} Avg_Convergence header mismatch")
            if d1.shape != d2.shape:
                problems.append(f"F{fid} Avg_Convergence shape mismatch: {d1.shape} vs {d2.shape}")
            else:
                # build mask: SFOA column (index 2) row 1 (index0) allowed to differ
                mask = np.ones(d1.shape, dtype=bool)
                mask[0, 2] = False
                diff = np.abs(d1 - d2)
                if not np.allclose(d1[mask], d2[mask], rtol=1e-9, atol=1e-6, equal_nan=True):
                    nd = int((diff[mask] > 1e-6).sum())
                    problems.append(f"F{fid} Avg_Convergence has {nd} non-SFOA-first-row cells differing from original >1e-6 (max|Δ|={diff[mask].max():.3e})")
                # (3) SFOA correction correctness
                rep_first = d2[0, 2]
                src_second = d1[1, 2]
                if abs(rep_first - src_second) > max(1e-3, 1e-9*abs(src_second)):
                    problems.append(f"F{fid} SFOA first-row correction error: repro={rep_first:.4e} should=orig_row2={src_second:.4e}")
                # (4) SFOA first row nonzero
                if rep_first == 0 or rep_first is None:
                    problems.append(f"F{fid} SFOA first row still 0")

            # (5) nfe
        if os.path.exists(rep_av):
            wb = openpyxl.load_workbook(rep_av, data_only=True)
            n = wb.active.max_row - 1
            wb.close()
            if n != EXPECT_NFE[dim]:
                problems.append(f"F{fid} nfe={n} expected {EXPECT_NFE[dim]}")

    return problems

if __name__ == "__main__":
    print("="*70)
    print("CEC2017 reproduction data vs original folder consistency check")
    print("Note: difference tolerance rtol=1e-9/atol=1e-6 (captures only real changes, ignores float serialization noise from openpyxl re-writing xlsx)")
    print("="*70)
    all_ok = True
    for dim in DIMS:
        probs = check_dim(dim)
        status = "PASS" if not probs else "FAIL"
        if probs: all_ok = False
        print(f"\n[Dim {dim}]  {status}  ({len(probs)} issues)")
        for p in probs[:20]:
            print("   -", p)
        if len(probs) > 20:
            print(f"   ... remaining {len(probs)-20} omitted")

    # (7) figure file completeness: each dimension figures/ should have 29 Boxplot + 29 Convergence TIFF
    print("\n" + "-"*70)
    print("Figure file completeness check (figures/*.tif)")
    print("-"*70)
    import glob
    for dim in DIMS:
        figdir = os.path.join(ROOT, f"CEC2017Dim{dim}_TIFF_MATLAB", "figures")
        box = glob.glob(os.path.join(figdir, "*Boxplot*.tif"))
        conv = glob.glob(os.path.join(figdir, "*Convergence.tif"))
        ok = (len(box)==29 and len(conv)==29)
        if not ok: all_ok = False
        print(f"  Dim {dim}: Boxplot={len(box)} Convergence={len(conv)}  {'OK' if ok else 'FAIL (expected 29 each)'}")

    print("\n" + "="*70)
    print("total result:", "ALL PASS" if all_ok else "HAS FAIL")
    print("="*70)
