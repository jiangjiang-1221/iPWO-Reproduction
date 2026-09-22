import re
from docx import Document
DOCX = r"C:/Users/JiangWen/Desktop/cn-20260826-review.docx"
d = Document(DOCX)
paras = d.paragraphs

print("===== key paragraphs (with keywords) =====")
keys = ["time window", "constraint", "augmented", "objective", "makespan", "cost Z", "Table", "Fig.", "Wilcoxon", "six types", "MTA", "VRPTW", "capacity", "31", "30", "6"]
for i, p in enumerate(paras):
    if any(k in p.text for k in keys):
        t = p.text.replace("\n", " ")
        print(f"[{i:3d}] {t[:110]}")

print("\n===== image rId + adjacent caption =====")
for i, p in enumerate(paras):
    xml = p._p.xml
    for rid in re.findall(r'embed="(rId\d+)"', xml):
        # search forward/backward for caption
        cap = ""
        for j in range(max(0,i-3), min(len(paras), i+4)):
            if "Fig." in paras[j].text:
                cap = paras[j].text.strip()[:80]; break
        print(f"{rid} @para{i}: {cap}")

print("\n===== Table 3 contents =====")
for ti, tbl in enumerate(d.tables):
    txt = [[c.text.strip() for c in r.cells] for r in tbl.rows]
    flat = " | ".join(" / ".join(r) for r in txt)
    if "iPWO" in flat or "Best" in flat:
        print(f"-- table {ti} --")
        for r in txt:
            print("  ", " | ".join(r))
