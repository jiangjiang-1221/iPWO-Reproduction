# -*- coding: utf-8 -*-
"""
Update cn-20260826-review.docx after the 6-UAV / 30-customer / TIME-WINDOW (VRPTW) re-run.

Extends the previous 6/30 updater with VRPTW-specific edits, all driven by the FRESH
comparison_table.csv produced by run_uav_delivery.m:
  * metric is now the constrained cost Z = makespan + time-window lateness penalty;
  * problem statement (para 76) lists 时间窗 as a constraint;
  * constraints subsection renumbered 五类 -> 六类, new 时间窗约束 paragraph inserted;
  * augmented-objective prose (para 100) notes the window penalty term;
  * statistics prose (para 111) regenerated from CSV, phrased as 约束代价 Z;
  * Table 3 rewritten from CSV;
  * conclusion (para 145) and robustness (para 121) stale numbers refreshed from CSV;
  * embedded Fig.5-8 replaced with new TIFFs, and a new Gantt figure (Fig.9) inserted.
A timestamped backup is made first.
"""
import os, re, csv, shutil, datetime, glob
from PIL import Image
import docx
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCX = r"C:/Users/江文/Desktop/cn-20260826-review.docx"
PKG  = r"D:/File/GPTTest/uav_delivery/uav_delivery_TIFF_MATLAB"
CSV  = os.path.join(PKG, "data", "comparison_table.csv")
FIGS = os.path.join(PKG, "figs")

assert os.path.exists(CSV), "comparison_table.csv not found - run not finished?"
assert os.path.exists(DOCX), "docx not found"

ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
bak = DOCX + f".bak_{ts}"
shutil.copy2(DOCX, bak)
print("backup ->", bak)

# ---------- parse CSV ----------
with open(CSV, encoding="utf-8-sig") as f:
    rows = [r for r in csv.reader(f) if r and any(c.strip() for c in r)]
header = [re.sub(r"[^a-z]", "", h.lower()) for h in rows[0]]
idx = {name: i for i, name in enumerate(header)}
def col(name):
    for k in idx:
        if name in k:
            return idx[k]
    raise KeyError(name)
ci_b, ci_med, ci_mean, ci_std = col("best"), col("median"), col("mean"), col("std")
ci_rank = [i for k, i in idx.items() if "rank" in k][0]
ci_p    = [i for k, i in idx.items() if "p" in k][0]
ci_algo = idx["algorithm"]
ALGOS = [r[ci_algo].strip() for r in rows[1:]]
data = {r[ci_algo].strip(): dict(best=float(r[ci_b]), median=float(r[ci_med]),
        mean=float(r[ci_mean]), std=float(r[ci_std]),
        rank=float(r[ci_rank]), p=r[ci_p].strip()) for r in rows[1:]}
print("algorithms:", ALGOS)

def f2(x):
    return f"{x:.2f}"

ipwo = data["iPWO"]; pwo = data["PWO"]
order = sorted(ALGOS, key=lambda a: data[a]["rank"])
rank_ipwo = order.index("iPWO") + 1
top = order[0]
min_std = min(data[a]["std"] for a in ALGOS)
conc = [a for a in ALGOS if data[a]["std"] <= 1.5 * min_std]
conc_sorted = sorted(conc, key=lambda a: data[a]["std"])
conc_range = f"{min(data[a]['std'] for a in conc):.2f}-{max(data[a]['std'] for a in conc):.2f}"
sig = [a for a in ALGOS if a != "iPWO" and data[a]["p"] not in ("-", "") and float(data[a]["p"]) < 0.05]
nonsig = [a for a in ALGOS if a != "iPWO" and a not in sig]
ratio = pwo["std"] / ipwo["std"]
ratio_s = f"{ratio:.0f}" if ratio >= 10 else f"{ratio:.1f}"
best_overall = min(ALGOS, key=lambda a: data[a]["best"])
is_ipwo_best = (best_overall == "iPWO")

def join_cn(names, sep="、"):
    if len(names) == 1: return names[0]
    if len(names) == 2: return f"{names[0]}{sep}{names[1]}"
    return "、".join(names[:-1]) + " 与 " + names[-1]

def pval_s(a):
    p = data[a]["p"]
    try:
        v = float(p)
        return f"{v:.1e}"
    except Exception:
        return p

sig_clause = join_cn([f"{a}（p={pval_s(a)}）" for a in sig]) if sig else "（无）"
nonsig_clause = "/".join(nonsig) if nonsig else "其余算法"
rank_phrase = (f"平均排名 Avg.Rank={ipwo['rank']:.2f}（全场最优）"
               if rank_ipwo == 1 else
               f"平均排名 Avg.Rank={ipwo['rank']:.2f}（位列第{rank_ipwo}，仅次于 {top} 的 {data[top]['rank']:.2f}）")
worst = order[-1]
best_clause = (f"并取得约束代价 Z 全场最低 {ipwo['best']:.2f} s"
               if is_ipwo_best else
               f"iPWO 的 Z 最优值为 {ipwo['best']:.2f} s（全场最低为 {best_overall} 的 {data[best_overall]['best']:.2f} s）")

Z_DESC = "约束代价 Z（完工时间 makespan + 时间窗迟到软惩罚）"
new_p111 = (
    f"51 次独立运行的统计指标显示：iPWO 取得 {Z_DESC} Best={ipwo['best']:.2f} s、"
    f"Median={ipwo['median']:.2f} s、Mean={ipwo['mean']:.2f} s、Std={ipwo['std']:.3f} s，{rank_phrase}；"
    f"{worst} 表现最差（Best={data[worst]['best']:.2f} s、Median={data[worst]['median']:.2f} s、"
    f"Mean={data[worst]['mean']:.2f} s、Std={data[worst]['std']:.3f} s、Avg.Rank={data[worst]['rank']:.2f}）。"
    f"如 Fig. 7 所示，箱线图（对数刻度）显示 iPWO 箱体接近最低且与 {'/'.join(conc_sorted)} 同属最集中一档"
    f"（Std {conc_range} s）；Wilcoxon 秩和检验（基于 51 次）表明 iPWO 在 p<0.05 水平显著优于 {sig_clause}，"
    f"而与 {nonsig_clause} 的差异未达显著（p>0.05）。值得注意的是，iPWO 相比原始 PWO 将代价离散度降低约 {ratio_s} 倍"
    f"（{ipwo['std']:.3f} s vs {pwo['std']:.3f} s），{best_clause}，说明其在强时间窗与容量约束下仍保持低方差与高可重复性。"
    f"上述统计量汇总于表 3（Table 3），时间窗可行性由 Fig. 9 甘特图直观验证。"
)

# ---------- load docx ----------
d = Document(DOCX)

def find_para(substr):
    for i, p in enumerate(d.paragraphs):
        if substr in p.text:
            return i
    return None

def set_text(i, txt):
    d.paragraphs[i].text = txt
    print(f"  set para {i}: {txt[:50]}...")

def insert_after(ref_idx, text=None, image_path=None, width_in=6.0, bold_cap=False):
    ref = d.paragraphs[ref_idx]._p
    if text is not None:
        np_ = OxmlElement('w:p'); ref.addnext(np_)
        cap = docx.text.paragraph.Paragraph(np_, d)
        r = cap.add_run(text)
        if bold_cap: r.bold = True
        ref = np_
    if image_path is not None:
        np2 = OxmlElement('w:p'); ref.addnext(np2)
        pp = docx.text.paragraph.Paragraph(np2, d)
        rr = pp.add_run()
        rr.add_picture(image_path, width=Inches(width_in))
    return ref_idx

# ---------- VRPTW prose edits ----------
i76 = find_para("目标是在满足容量、能耗与空域安全等约束")
if i76 is not None:
    set_text(i76, d.paragraphs[i76].text.replace(
        "在满足容量、能耗与空域安全等约束下最",
        "在满足容量、能耗、时间窗与空域安全等约束下最"))
else:
    print("  [skip] para76 not found")

i3 = find_para("受容量、障碍、冲突、运动学与能耗等多重硬约束")
if i3 is not None:
    set_text(i3, d.paragraphs[i3].text.replace(
        "受容量、障碍、冲突、运动学与能耗等多重硬约束",
        "受容量、时间窗、障碍、冲突、运动学与能耗等多重硬约束"))
else:
    print("  [skip] para3 not found")

i88 = find_para("在优化过程中须同时满足以下五类工程约束")
if i88 is not None:
    set_text(i88, d.paragraphs[i88].text.replace("五类", "六类"))
else:
    print("  [skip] para88 not found")

i93 = find_para("运动学约束：相邻航段最大爬升角不超过 90")
if i93 is not None:
    tw_txt = ("时间窗约束：每个客户 j 被赋予服务时间窗 [e_j, l_j]（早到需在窗前等待、晚到产生软惩罚 "
              "ω_t·(t−l_j)）；该约束将可行性边界引入“分配—排序”空间，仅当任务分配与访问顺序同时满足时间窗时方可获得低代价解，"
              "从而凸显算法在强约束下的寻优能力。")
    insert_after(i93, text=tw_txt)
    print("  inserted 时间窗约束 paragraph after", i93)
else:
    print("  [skip] para93 not found")

i100 = find_para("常数 100 使任一约束违反带来的目标增量远超可行解 makespan")
if i100 is not None:
    set_text(i100, d.paragraphs[i100].text.replace(
        "常数 100 使任一约束违反带来的目标增量远超可行解 makespan（约 165 s），从而引导种群优先消除不可行、再压缩完工时间。",
        "常数 100 使容量/障碍/冲突/运动学等硬约束违反带来的目标增量远超可行解代价（约 165 s）；时间窗迟到惩罚以权重 "
        "ω_t=2.0 计入增广目标，使约束代价 Z = makespan + 时间窗迟到在可行解处退化为真实完工时间，从而引导种群优先消除不可行、再压缩约束代价 Z。"))
else:
    print("  [skip] para100 not found")

# ---------- statistics (para 111) ----------
i111 = find_para("51 次独立运行的统计指标显示")
if i111 is not None:
    set_text(i111, new_p111)
    print("  regenerated para", i111)
else:
    print("  [skip] para111 not found")

# ---------- conclusion (para 145) ----------
i145 = find_para("全场最低最优完工时间")
if i145 is not None:
    t = d.paragraphs[i145].text
    t = re.sub(r"全场最低最优完工时间 [\d.]+ s、平均 [\d.]+ s、标准差 [\d.]+ s，Avg\.Rank=[\d.]+",
               f"约束代价 Z 最低 {ipwo['best']:.2f} s、平均 {ipwo['mean']:.2f} s、标准差 {ipwo['std']:.3f} s，Avg.Rank={ipwo['rank']:.2f}",
               t)
    set_text(i145, t)
    print("  updated conclusion para", i145)
else:
    print("  [skip] para145 not found")

# ---------- robustness (para 121) ----------
i121 = find_para("iPWO 的标准差（0.884 s）")
if i121 is not None:
    t = d.paragraphs[i121].text
    t = t.replace("0.884 s", f"{ipwo['std']:.3f} s")
    t = re.sub(r"PWO 的 [\d.]+", f"PWO 的 {pwo['std']:.3f}", t)
    set_text(i121, t)
    print("  updated robustness para", i121)
else:
    print("  [skip] para121 not found")

# ---------- Table 3 rewrite ----------
def clean_p(p):
    try:
        v = float(p)
        return f"{v:.1e}" if v < 1e-3 else f"{v:.1e}"
    except Exception:
        return p if p not in ("-", "") else "—"

tbl = d.tables[3]
labels = ["Best (s)", "Median (s)", "Mean (s)", "Std (s)", "Avg.Rank", "p-value"]
tbl.rows[0].cells[0].text = "Algorithm"
for k, lab in enumerate(labels):
    tbl.rows[0].cells[1 + k].text = lab
for r, a in enumerate(ALGOS, start=1):
    row = tbl.rows[r]
    row.cells[0].text = a
    row.cells[1].text = f"{data[a]['best']:.2f}"
    row.cells[2].text = f"{data[a]['median']:.2f}"
    row.cells[3].text = f"{data[a]['mean']:.2f}"
    row.cells[4].text = f"{data[a]['std']:.3f}"
    row.cells[5].text = f"{data[a]['rank']:.2f}"
    row.cells[6].text = clean_p(data[a]["p"])
print("rewrote Table 3")

# ---------- replace embedded Fig.5-8 + insert Gantt ----------
def tif_to_png(tif_path, max_w=1800):
    im = Image.open(tif_path)
    if im.mode != "RGB":
        im = im.convert("RGB")
    if im.width > max_w:
        h = int(im.height * max_w / im.width)
        im = im.resize((max_w, h), Image.LANCZOS)
    out = tif_path + ".embed.png"
    im.save(out, "PNG", optimize=True)
    return out

mapping = {"Fig. 5": "fig1_scene", "Fig. 6": "fig4_convergence",
           "Fig. 7": "fig5_boxplot", "Fig. 8": "fig6_conflict"}
rels = d.part.rels
embed_rids = {}
for i, para in enumerate(d.paragraphs):
    xml = para._p.xml
    for rid in re.findall(r'embed="(rId\d+)"', xml):
        for cap_frag in mapping:
            for j in range(max(0, i - 3), min(len(d.paragraphs), i + 4)):
                if cap_frag in d.paragraphs[j].text:
                    embed_rids[rid] = cap_frag
                    break
print("figure rIds:", embed_rids)
for rid, cap in embed_rids.items():
    base = mapping[cap]
    tif = glob.glob(os.path.join(FIGS, base + ".tif")) or glob.glob(os.path.join(FIGS, base + ".png"))
    if not tif:
        print("  WARN: missing", base); continue
    png = tif_to_png(tif[0])
    with open(png, "rb") as fh:
        rels[rid].blob = fh.read()
    print(f"  replaced {rid} ({cap}) <- {os.path.basename(tif[0])}")

# insert Gantt (Fig.9) after the statistics paragraph
gantt = glob.glob(os.path.join(FIGS, "fig9_gantt.tif"))
if gantt and i111 is not None:
    gp = tif_to_png(gantt[0])
    insert_after(i111, text="Fig. 9 时间窗可行性甘特图（最优 iPWO 解的到达时刻与服务时段，浅蓝阴影为各客户时间窗 [e_j, l_j]）",
                 image_path=gp, width_in=6.0, bold_cap=True)
    print("  inserted Gantt figure (Fig.9)")

d.save(DOCX)
print("docx saved:", DOCX)
print("\n--- new para 111 ---\n", new_p111)
