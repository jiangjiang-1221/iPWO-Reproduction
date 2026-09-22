#!/usr/env python
# -*- coding: utf-8 -*-
"""One-shot: re-embed the two corrected figures (Fig.6 mean convergence rId97 + Fig.5 boxplot rId98)
# NOTE: this utility edits the Chinese manuscript copy. The Chinese literals below are
# match/replacement strings for that document and are kept verbatim on purpose.
into the original paper cn-20260826-review.docx and fix the Fig.6 description/caption text so they
match Table 3. Back up the original first, then overwrite in place."""
import os, glob, shutil
from datetime import datetime
from PIL import Image
import docx

SCRIPT_DIR = r"D:\File\GPTTest\uav_delivery\uav_delivery_TIFF_MATLAB"
FIGS = os.path.join(SCRIPT_DIR, "figs")
SRC = r"C:/Users/USER/Desktop/cn-20260826-review.docx"
BAK = SRC + ".bak_before_figfix_" + datetime.now().strftime("%Y%m%d_%H%M%S")

def tif_to_png(tif, max_w=1800):
    im = Image.open(tif)
    if im.mode != "RGB":
        im = im.convert("RGB")
    if im.width > max_w:
        h = int(im.height * max_w / im.width)
        im = im.resize((max_w, h), Image.LANCZOS)
    out = tif + ".embed.png"
    im.save(out, "PNG", optimize=True)
    return out

# 0) backup original
shutil.copy2(SRC, BAK)
print("backup ->", BAK)

d = docx.Document(SRC)
rels = d.part.rels

# 1) Fig.6 mean convergence (rId97)
tif6 = glob.glob(os.path.join(FIGS, "fig4_convergence.tif"))
assert tif6, "fig4_convergence.tif not found"
png6 = tif_to_png(tif6[0])
with open(png6, "rb") as fh:
    rels["rId97"].target_part._blob = fh.read()
print("re-embedded Fig.6 (rId97) <-", os.path.basename(png6))

# 2) Fig.5 boxplot (rId98)
tif5 = glob.glob(os.path.join(FIGS, "fig5_boxplot.tif"))
assert tif5, "fig5_boxplot.tif not found"
png5 = tif_to_png(tif5[0])
with open(png5, "rb") as fh:
    rels["rId98"].target_part._blob = fh.read()
print("re-embedded boxplot (rId98) <-", os.path.basename(png5))

# 3) fix Fig.6 description paragraph
for p in d.paragraphs:
    if "如Fig. 6所示" in p.text:  # Chinese docx caption marker; kept verbatim
        t = p.text
        t = t.replace(
            "各算法含惩罚项的目标函数收敛曲线（横轴为迭代次数，纵轴为含惩罚项的增广目标函数 F 的取值，采用对数刻度，越低越优）",
            "各算法约束代价 Z 的均值收敛曲线（横轴为迭代次数，纵轴为约束代价 Z = 完工时间 + 时间窗迟到，与表 3 同一指标，采用对数刻度，越低越优）")  # Chinese docx text kept verbatim
        t = t.replace(
            "每条曲线代表一种算法在51次独立运行中的最优结果；纵坐标为增广目标函数F（图中标记为F），即完工时间加上惩罚项。",
            "每条曲线为 51 次独立运行的均值（best-so-far Z 的均值），纵坐标即约束代价 Z = 完工时间 + 时间窗迟到，与表 3 完全一致；")  # Chinese docx text kept verbatim
        p.text = t
        print("updated description para")

# 4) fix Fig.6 caption
for p in d.paragraphs:
    if p.text.strip().startswith("Fig. 6 Convergence curves"):
        p.text = "Fig. 6 Mean convergence curves of the constrained cost Z = makespan + time-window lateness over 51 runs (log scale, lower is better)."
        print("updated caption para")

# 5) overwrite original in place
d.save(SRC)
print("OVERWRITTEN_ORIGINAL_OK ->", SRC)
