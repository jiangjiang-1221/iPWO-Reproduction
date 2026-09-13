#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""重嵌新的均值收敛图 Fig.6 到论文，并更新描述段/图注（与 Table 3 同指标 Z）。
依赖：figs/fig4_convergence.tif 必须已由 run_uav_delivery.m 重新生成。
"""
import os, re, shutil, glob, datetime
from PIL import Image
import docx
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
FIGS = os.path.join(SCRIPT_DIR, "figs")
DOCX = r"C:/Users/江文/Desktop/cn-20260826-review.docx"

# 备份
stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
bak = DOCX + f".bak_fig6_{stamp}"
shutil.copy2(DOCX, bak)
print("backup:", bak)

def tif_to_png(tif_path, max_w=1800):
    im = Image.open(tif_path)
    if im.mode != "RGB":
        im = im.convert("RGB")
    if im.width > max_w:
        h = int(im.height * max_w / im.width)
        im = im.resize((max_w, h), Image.LANCZOS)
    out = tif_path + ".embed.png"
    im.save(out, "PNG", optimize=True)
    return out

d = docx.Document(DOCX)
rels = d.part.rels

# ---- 1) 重嵌 Fig.6 (rId97, 已校验) ----
tif = glob.glob(os.path.join(FIGS, "fig4_convergence.tif"))
assert tif, "fig4_convergence.tif 未找到，请先跑完 run_uav_delivery.m"
png = tif_to_png(tif[0])
with open(png, "rb") as fh:
    rels["rId97"].target_part._blob = fh.read()
print("re-embedded Fig.6 (rId97) <-", os.path.basename(png))

# ---- 2) 改描述段（含 "如Fig. 6所示"） ----
for p in d.paragraphs:
    if "如Fig. 6所示" in p.text:
        t = p.text
        t = t.replace(
            "各算法含惩罚项的目标函数收敛曲线（横轴为迭代次数，纵轴为含惩罚项的增广目标函数 F 的取值，采用对数刻度，越低越优）",
            "各算法约束代价 Z 的均值收敛曲线（横轴为迭代次数，纵轴为约束代价 Z = 完工时间 + 时间窗迟到，与表 3 同一指标，采用对数刻度，越低越优）")
        t = t.replace(
            "每条曲线代表一种算法在51次独立运行中的最优结果；纵坐标为增广目标函数F（图中标记为F），即完工时间加上惩罚项。",
            "每条曲线为 51 次独立运行的均值（best-so-far Z 的均值），纵坐标即约束代价 Z = 完工时间 + 时间窗迟到，与表 3 完全一致；")
        p.text = t
        print("updated description para")

# ---- 3) 改图注 (Fig. 6 ...) ----
for p in d.paragraphs:
    if p.text.strip().startswith("Fig. 6 Convergence curves"):
        p.text = "Fig. 6 Mean convergence curves of the constrained cost Z = makespan + time-window lateness over 51 runs (log scale, lower is better)."
        print("updated caption para")

d.save(DOCX)
print("docx saved:", DOCX)
